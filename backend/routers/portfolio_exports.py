"""Portfolio export routes: GET /portfolio/export/excel/{user_id} and /portfolio/export/pdf/{user_id}.

Split out of routers/portfolio.py (ARCH-14); paths, handler names and tags are unchanged.
"""
import logging
from datetime import datetime, date, timedelta, timezone
from fastapi import APIRouter, Depends, HTTPException, Header
from core.security import require_user_match
from fastapi.responses import StreamingResponse
from helpers import (
    get_session, _to_ns, _yf_price, User, Portfolio, AnalysisHistory,
    SignalEvent, format_indian_currency,
)
import importlib.util as _importlib_util

router = APIRouter(tags=["Portfolio"])


FPDF2_AVAILABLE = _importlib_util.find_spec("fpdf") is not None  # fpdf2 installs the `fpdf` package


if FPDF2_AVAILABLE:
    from fpdf import FPDF


try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font as XLFont, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


# ---------------------------------------------------------------------------
# 12. GET /portfolio/export/excel/{user_id}
# ---------------------------------------------------------------------------
@router.get("/portfolio/export/excel/{user_id}")
def export_portfolio_excel(user_id: int, _caller: int = Depends(require_user_match)):
    if not OPENPYXL_AVAILABLE:
        raise HTTPException(status_code=501, detail="openpyxl not installed on server")
    try:
        with get_session() as session:
            user = session.get(User, user_id)
            if not user:
                raise HTTPException(status_code=404, detail=f"User {user_id} not found")
            holdings = session.query(Portfolio).filter_by(user_id=user_id, group_name="holdings").all()

            last_signals: dict[str, str] = {}
            last_confidence: dict[str, float | None] = {}
            for h in holdings:
                row = session.query(AnalysisHistory)\
                    .filter(AnalysisHistory.ticker == h.ticker)\
                    .order_by(AnalysisHistory.created_at.desc()).first()
                if row:
                    last_signals[h.ticker] = row.signal
                    last_confidence[h.ticker] = row.confidence_pct

            # The caller's own analyses only; never other users' research.
            history_rows = session.query(AnalysisHistory)\
                .filter(AnalysisHistory.user_id == user_id)\
                .order_by(AnalysisHistory.created_at.desc()).limit(100).all()
            # Detach data before session closes
            history_data = [(r.analysis_date, r.ticker, r.signal, r.confidence_pct,
                             r.target_short, r.target_medium, r.target_long) for r in history_rows]
            holdings_data = [(h.ticker, h.company_name, h.shares, h.buy_price,
                              h.notes or "") for h in holdings]

        wb = openpyxl.Workbook()

        # ── Sheet 1: Portfolio ──────────────────────────────────────
        ws1 = wb.active
        ws1.title = "Portfolio"

        headers = ["Ticker", "Company", "Shares", "Buy Price (Rs)", "Current Price (Rs)",
                   "Current Value (Rs)", "Invested (Rs)", "P&L (Rs)", "P&L %",
                   "Signal", "Confidence %", "Notes"]

        header_fill = PatternFill(start_color="2C2C2C", end_color="2C2C2C", fill_type="solid")
        header_font = XLFont(color="F5F0E8", bold=True, size=11)

        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[1].height = 22

        col_widths = [12, 32, 10, 18, 20, 20, 18, 16, 10, 20, 14, 40]
        for i, w in enumerate(col_widths, 1):
            ws1.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

        green_fill = PatternFill(start_color="D6F5E3", end_color="D6F5E3", fill_type="solid")
        red_fill   = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")

        for row_num, (ticker, company, shares, buy_price, notes) in enumerate(holdings_data, 2):
            ns = _to_ns(ticker)
            current_price, _ = _yf_price(ns)
            if not current_price:
                current_price = 0.0
            invested       = round(shares * buy_price, 2)
            current_value  = round(shares * current_price, 2)
            pnl            = round(current_value - invested, 2)
            pnl_pct        = round((pnl / invested * 100) if invested > 0 else 0, 2)
            signal         = last_signals.get(ticker, "—")
            confidence     = last_confidence.get(ticker, "—")

            row_data = [ticker, company, shares, buy_price, current_price,
                        current_value, invested, pnl, pnl_pct, signal, confidence, notes]
            for col, value in enumerate(row_data, 1):
                cell = ws1.cell(row=row_num, column=col, value=value)
                cell.alignment = Alignment(
                    horizontal="center" if col in (1, 3, 10, 11) else "right"
                )
                if col in (8, 9):
                    cell.fill = green_fill if pnl >= 0 else red_fill

        # ── Sheet 2: Analysis History ───────────────────────────────
        ws2 = wb.create_sheet("Analysis History")
        hist_headers = ["Date", "Ticker", "Signal", "Confidence %",
                        "Target 3M (Rs)", "Target 12M (Rs)", "Target 3Y (Rs)"]
        for col, h in enumerate(hist_headers, 1):
            cell = ws2.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for w, col in zip([14, 14, 22, 14, 18, 18, 18], range(1, 8)):
            ws2.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

        for idx, (adate, aticker, signal, conf, t3m, t12m, t36m) in enumerate(history_data, 2):
            ws2.cell(row=idx, column=1, value=adate)
            ws2.cell(row=idx, column=2, value=aticker)
            ws2.cell(row=idx, column=3, value=signal)
            ws2.cell(row=idx, column=4, value=conf)
            ws2.cell(row=idx, column=5, value=t3m)
            ws2.cell(row=idx, column=6, value=t12m)
            ws2.cell(row=idx, column=7, value=t36m)

        import io
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        today = date.today().strftime("%Y%m%d")
        filename = f"nivesh_portfolio_{today}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail="Internal server error")


# ---------------------------------------------------------------------------
# 23. GET /portfolio/export/pdf/{user_id}
# ---------------------------------------------------------------------------
@router.get("/portfolio/export/pdf/{user_id}")
def export_portfolio_pdf(user_id: int, _caller: int = Depends(require_user_match)):
    """Generate a deterministic, factual portfolio summary PDF (no LLM call, no advice)."""
    if not FPDF2_AVAILABLE:
        raise HTTPException(status_code=503, detail="PDF export unavailable")
    try:
        import io

        with get_session() as session:
            user = session.get(User, user_id)
            if not user:
                raise HTTPException(status_code=404, detail=f"User {user_id} not found")
            holdings = session.query(Portfolio).filter_by(user_id=user_id, group_name="holdings").all()
            holdings_snap = [(h.ticker, h.company_name, h.shares, h.buy_price) for h in holdings]
            tickers = [t for t, _, _, _ in holdings_snap]

            # Watch items: factual flags only (pledging is a company-level fact).
            red_flag_list = []
            cutoff30 = datetime.now(timezone.utc) - timedelta(days=30)
            bf_map = {
                f.ticker for f in session.query(SignalEvent).filter(
                    SignalEvent.created_at >= cutoff30,
                    SignalEvent.ticker.in_(tickers),
                    SignalEvent.to_signal.in_(["BEARISH", "CAUTIOUSLY BEARISH"]),
                ).all()
            }
            for ticker in tickers:
                pledged_row = session.query(AnalysisHistory.promoter_pledged_pct).filter(
                    AnalysisHistory.ticker == ticker,
                    AnalysisHistory.promoter_pledged_pct.isnot(None),
                ).order_by(AnalysisHistory.created_at.desc()).first()
                pledged = pledged_row[0] if pledged_row else None
                if pledged is not None and pledged > 30:
                    red_flag_list.append(f"{ticker}: Promoter pledging {pledged:.1f}%")
                if ticker in bf_map:
                    red_flag_list.append(f"{ticker}: Signal turned bearish in the last 30 days")

        # Holdings with no live price show N/A and are left out of BOTH totals,
        # so a failed fetch never reads as a loss.
        rows = []
        total_invested = 0.0
        total_current = 0.0
        unpriced = 0
        for ticker, _company, shares, buy_price in holdings_snap:
            current_price, _ = _yf_price(_to_ns(ticker))
            invested = shares * buy_price
            if not (current_price and current_price > 0):  # None, 0 or NaN
                unpriced += 1
                rows.append((ticker, shares, invested, None, None))
                continue
            current_val = shares * current_price
            pnl_pct = (current_val - invested) / invested * 100 if invested > 0 else None
            total_invested += invested
            total_current += current_val
            rows.append((ticker, shares, invested, current_val, pnl_pct))

        def _rs(value):
            # Core Helvetica font is latin-1 only, so no rupee glyph.
            return format_indian_currency(value, symbol="Rs. ")

        try:
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Helvetica", style="B", size=16)
            pdf.cell(0, 10, "Portfolio Summary", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=9)
            pdf.cell(0, 6, f"Generated on {date.today().strftime('%d %b %Y')}", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)

            widths = (40, 25, 42, 42, 31)
            pdf.set_font("Helvetica", style="B", size=10)
            for w, head in zip(widths, ("Ticker", "Shares", "Invested", "Current value", "P&L %")):
                pdf.cell(w, 8, head, border=1)
            pdf.ln(8)
            pdf.set_font("Helvetica", size=10)
            for ticker, shares, invested, current_val, pnl_pct in rows:
                cells = (
                    ticker, f"{shares:.4f}".rstrip("0").rstrip("."), _rs(invested),
                    "N/A" if current_val is None else _rs(current_val),
                    "N/A" if pnl_pct is None else f"{pnl_pct:+.1f}%",
                )
                for w, text in zip(widths, cells):
                    pdf.cell(w, 7, text, border=1)
                pdf.ln(7)
            if not rows:
                pdf.cell(0, 7, "No holdings found.", new_x="LMARGIN", new_y="NEXT")

            pdf.ln(4)
            pdf.set_font("Helvetica", style="B", size=11)
            pdf.cell(0, 7, f"Total invested: {_rs(total_invested)}", new_x="LMARGIN", new_y="NEXT")
            pnl_txt = f"  ({(total_current - total_invested) / total_invested * 100:+.1f}%)" if total_invested > 0 else ""
            pdf.cell(0, 7, f"Current value: {_rs(total_current)}{pnl_txt}", new_x="LMARGIN", new_y="NEXT")
            if unpriced:
                pdf.set_font("Helvetica", style="I", size=9)
                pdf.cell(0, 6, f"Totals exclude {unpriced} holding(s) with no live price.", new_x="LMARGIN", new_y="NEXT")
                pdf.set_font("Helvetica", style="B", size=11)

            pdf.ln(4)
            pdf.cell(0, 7, "Watch items", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=10)
            for flag in red_flag_list or ["None detected"]:
                pdf.cell(0, 6, f"- {flag}", new_x="LMARGIN", new_y="NEXT")

            pdf.ln(6)
            pdf.set_font("Helvetica", style="I", size=8)
            pdf.multi_cell(0, 5, "Informational summary only. Not investment advice.")
            pdf_bytes = bytes(pdf.output())
        except Exception as _pdfe:
            logging.error(f"[PDF-EXPORT] fpdf2 render failed: {_pdfe}", exc_info=True)
            raise HTTPException(status_code=500, detail="PDF export failed")

        today_str = date.today().strftime("%Y%m%d")
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=portfolio_summary_{today_str}.pdf"},
        )

    except HTTPException:
        raise
    except Exception:
        logging.exception("[PDF-EXPORT] Top-level error")
        raise HTTPException(status_code=503, detail="Portfolio review unavailable. Please try again later.")
